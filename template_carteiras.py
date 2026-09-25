# -*- coding: utf-8 -*-
"""template_carteiras.py — leitura do TemplateCarteiras.xlsx do ControleCargas.

[2026-09-03, pedido do usuário] A ferramenta de Publicar Agrupamentos deste
projeto não deve oferecer para publicação nenhum agrupamento que tenha, entre
seus membros, uma carteira marcada com "Deve Publicar" = Não no cadastro
TemplateCarteiras.xlsx — cadastro mantido no projeto irmão ControleCargas
(`registry.py`/`snapshot_builder.py` de lá leem a mesma coluna, mesma regra:
`devePublicar.strip().casefold() == "sim"`). Decisão do usuário: ler o MESMO
arquivo Excel direto (sem duplicar cópia nem expor rota nova no
ControleCargas).

[2026-09-24, pedido do usuário] Os painéis "Painel de Controle - Template
Carteiras" e "... (Somente SLA)" usam este mesmo cadastro como universo de
carteiras (ver wallet_scope.py) — por isso a leitura passou a devolver o
cadastro inteiro (nome, instituição, defasagem, agrupamentos), não só o set
de bloqueadas.

[CORRIGIDO 2026-09-24] O caminho padrão apontava para a pasta `data/` do
CLONE Git do ControleCargas (`../ControleCargas/prototype/data`) — cópia
local que não recebe as edições do time (estava parada em 24/08 enquanto a
pasta compartilhada já tinha as carteiras novas). O arquivo vivo é o da
biblioteca do SharePoint sincronizada pelo OneDrive
(`~/Beehus Tecnologia Ltda/Beehus Tecnologia Ltda - Documentos/SWAT/
ControleCargas/prototype/data`) — mesma pasta que o ControleCargas resolve
em `utils/caminhos.py::resolver_data_dir`. O clone fica só como último
recurso, com aviso no log.

Somente leitura; este módulo nunca escreve no arquivo. Cache por mtime do
arquivo (não por TTL fixo): como é editado manualmente por humanos com pouca
frequência, vale mais reagir à mudança real do que esperar um TTL vencer.
"""
import logging
import os
import threading
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

_NOME_ARQUIVO = "TemplateCarteiras.xlsx"
# Trecho estável do caminho dentro da biblioteca compartilhada — o que muda de
# uma máquina para outra é só a pasta-raiz da biblioteca (ver candidatos).
_SUFIXO_BIBLIOTECA = Path("SWAT") / "ControleCargas" / "prototype" / "data" / _NOME_ARQUIVO
# Cópia do clone Git irmão (mesma pasta "Projeto - Servidor") — último recurso.
_CAMINHO_CLONE = Path(os.path.dirname(os.path.abspath(__file__))).parent / \
    "ControleCargas" / "prototype" / "data" / _NOME_ARQUIVO

# Colunas fixas do TemplateCarteiras (A-Q) — mesmo contrato de
# ControleCargas/prototype/registry.py::_COLUNAS_TEMPLATE (índice 0-based).
_COL_NOME = 0            # A  Nome Carteira
_COL_WALLET_ID = 1       # B  WalletID
_COL_INSTITUICAO = 2     # C  Instituição
_COL_DEVE_PUBLICAR = 4   # E  Deve Publicar
_COL_AGRUPAMENTOS = 6    # G  Agrupamentos Indexados
_COL_PERIODICIDADE = 7   # H  Periodicidade
_COL_DEFASAGEM = 8       # I  Defasagem

_lock = threading.Lock()
_cache = {"caminho": None, "mtime": None, "carteiras": {}}


def _candidatos():
    """Contexto:
    Lista, em ordem de preferência, os caminhos possíveis do
    TemplateCarteiras.xlsx nesta máquina — insumo de caminho_template().

    Pseudocódigo:
      1. Variável de ambiente TEMPLATE_CARTEIRAS_PATH (override explícito).
      2. Biblioteca do SharePoint sincronizada na pasta do usuário, nas duas
         grafias que o OneDrive usa ("- Documentos" / "- Documents").
      3. A mesma biblioteca dentro do OneDrive corporativo pessoal
         (OneDriveCommercial/OneDrive e a pasta-pai delas).
      4. O clone Git irmão (cópia local, pode estar desatualizada) — antes
         dele, caminho_template() tenta a busca por variantes.
    """
    caminhos = []
    override = os.environ.get("TEMPLATE_CARTEIRAS_PATH")
    if override:
        caminhos.append(Path(override))
    home = Path.home()
    for biblioteca in ("Beehus Tecnologia Ltda - Documentos",
                       "Beehus Tecnologia Ltda - Documents"):
        caminhos.append(home / "Beehus Tecnologia Ltda" / biblioteca / _SUFIXO_BIBLIOTECA)
    for variavel in ("OneDriveCommercial", "OneDrive"):
        raiz = os.environ.get(variavel)
        if raiz:
            caminhos.append(Path(raiz) / _SUFIXO_BIBLIOTECA)
            caminhos.append(Path(raiz).parent / "Beehus Tecnologia Ltda"
                            / "Beehus Tecnologia Ltda - Documentos" / _SUFIXO_BIBLIOTECA)
    caminhos.append(_CAMINHO_CLONE)
    return caminhos


# Marca das pastas onde o OneDrive sincroniza a biblioteca do time.
_MARCA_PASTA_BIBLIOTECA = "beehus"
# Só o achado é memorizado: não achando, a próxima chamada procura de novo
# (a pessoa pode sincronizar a biblioteca com o app no ar).
_variante_encontrada = {"caminho": None}


def _procurar_variante():
    """Contexto:
    [CORRIGIDO 2026-09-25, relato do usuário: "uma colega de time está
    rodando, em Painel - Template (SLA) não está aparecendo as transações"]
    Os caminhos fixos de _candidatos() são os nomes que o OneDrive dá à
    biblioteca NESTA máquina; em outra eles mudam (Windows em inglês,
    biblioteca no OneDrive pessoal "OneDrive - Beehus .../SWAT/...", pasta
    renomeada). Sem achar o arquivo, carteiras_template() devolvia {}, o
    escopo dos painéis Template virava um conjunto VAZIO e toda transação
    era filtrada — sem erro na tela. Mesma busca que o ControleCargas faz em
    utils/caminhos.py::_procurar_data_dir_compartilhado (achado de 22/09).

    Pseudocódigo:
      1. Bases: pasta do usuário e as raízes OneDrive/OneDriveCommercial
         (e os pais delas — a biblioteca fica AO LADO do OneDrive pessoal).
      2. Em cada base, só as subpastas com "beehus" no nome (listagem curta,
         nunca varre o disco — pasta "só na nuvem" custaria download).
      3. Testa <pasta>/SWAT/... e, não achando, 1 nível abaixo
         (<pasta>/<biblioteca>/SWAT/...). Primeiro arquivo que existir vence.
    """
    if _variante_encontrada["caminho"] is not None:
        return _variante_encontrada["caminho"]
    bases = [Path.home()]
    for variavel in ("OneDriveCommercial", "OneDrive"):
        raiz = os.environ.get(variavel)
        if raiz:
            bases.extend([Path(raiz), Path(raiz).parent])
    vistas = set()
    for base in bases:
        chave = str(base).casefold()
        if chave in vistas:
            continue
        vistas.add(chave)
        try:
            pastas = [p for p in base.iterdir()
                      if p.is_dir() and _MARCA_PASTA_BIBLIOTECA in p.name.casefold()]
        except OSError:
            continue
        for pasta in pastas:
            try:
                niveis = [pasta] + [p for p in pasta.iterdir() if p.is_dir()]
            except OSError:
                niveis = [pasta]
            for nivel in niveis:
                caminho = nivel / _SUFIXO_BIBLIOTECA
                try:
                    if caminho.is_file():
                        _variante_encontrada["caminho"] = caminho
                        return caminho
                except OSError:
                    continue
    return None


def caminho_template():
    """Primeiro candidato de _candidatos() que existe em disco, ou None. Antes
    do clone (último recurso) tenta a busca por variantes de _procurar_variante()."""
    for caminho in _candidatos():
        if caminho == _CAMINHO_CLONE:
            variante = _procurar_variante()
            if variante is not None:
                return str(variante)
        if caminho.is_file():
            if caminho == _CAMINHO_CLONE:
                logger.warning(
                    "TemplateCarteiras.xlsx da pasta compartilhada não encontrado — "
                    "usando a cópia do clone Git (%s), que pode estar desatualizada.",
                    caminho)
            return str(caminho)
    return None


def _texto(valor):
    return str(valor if valor is not None else "").strip()


def _interpretar_defasagem(valor):
    """Contexto: texto da coluna "Defasagem" ("D-3", "M", vazio) -> nº de dias
    úteis. [2026-09-24, decisão do usuário] "M" (mensal) e vazio viram 1 —
    no painel Somente SLA a carteira mensal é tratada como D-1. (Diverge de
    propósito do ControleCargas, que usa 0 para "M".)

    Pseudocódigo:
      1. "D-<n>" com n inteiro -> n.
      2. Qualquer outra coisa ("M", vazio, texto livre) -> 1.
    """
    texto = _texto(valor).upper()
    if texto.startswith("D-"):
        try:
            return int(texto[2:])
        except ValueError:
            return 1
    return 1


def _interpretar_agrupamentos(valor):
    """"id1;id2" -> ["id1", "id2"]; "Não"/vazio -> []. Mesma regra de
    ControleCargas/prototype/registry.py::_interpretar_agrupamentos."""
    texto = _texto(valor)
    if not texto or texto.casefold() in ("não", "nao"):
        return []
    return [g.strip() for g in texto.split(";") if g.strip()]


def _ler_carteiras(caminho):
    """Contexto:
    Lê o TemplateCarteiras.xlsx e devolve {walletId: info}. Chamada por
    carteiras_template() só quando o arquivo (caminho ou mtime) muda.

    Pseudocódigo:
      1. Abre em read_only + iter_rows (varredura sequencial; o workbook é
         sempre fechado no finally para não prender o arquivo do Excel).
      2. Pula o cabeçalho e as linhas sem WalletID (cauda vazia).
      3. Monta o info da linha: nome, instituição, deve publicar, lista de
         agrupamentos, periodicidade, defasagem original e em dias úteis.
    """
    workbook = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    try:
        planilha = workbook[workbook.sheetnames[0]]
        carteiras = {}
        for linha in planilha.iter_rows(min_row=2, values_only=True):
            if not linha or len(linha) <= _COL_DEFASAGEM:
                continue
            wallet_id = _texto(linha[_COL_WALLET_ID])
            if not wallet_id:
                continue
            carteiras[wallet_id] = {
                "walletId": wallet_id,
                "nome": _texto(linha[_COL_NOME]),
                "instituicao": _texto(linha[_COL_INSTITUICAO]),
                "devePublicar": _texto(linha[_COL_DEVE_PUBLICAR]).casefold() == "sim",
                "agrupamentos": _interpretar_agrupamentos(linha[_COL_AGRUPAMENTOS]),
                "periodicidade": _texto(linha[_COL_PERIODICIDADE]).upper() or "D",
                "defasagem": _texto(linha[_COL_DEFASAGEM]).upper(),
                "defasagemDu": _interpretar_defasagem(linha[_COL_DEFASAGEM]),
            }
        return carteiras
    finally:
        workbook.close()


def carteiras_template():
    """Contexto:
    {walletId: info} de TODAS as carteiras do TemplateCarteiras.xlsx —
    universo dos painéis Template Carteiras (wallet_scope.py) e base do
    bloqueio de publicação. Cacheado pelo (caminho, mtime) do arquivo.
    Não mutar o dict devolvido.

    Defensivo: arquivo ausente/inacessível -> {} (loga aviso); falha de
    leitura (arquivo sendo salvo no instante) -> mantém o último resultado
    válido em cache. Nunca lança exceção.

    Pseudocódigo:
      1. Resolve o caminho; sem arquivo -> {}.
      2. (caminho, mtime) iguais ao cache -> devolve o cache.
      3. Senão relê o arquivo e atualiza o cache.
    """
    caminho = caminho_template()
    if not caminho:
        logger.warning("TemplateCarteiras.xlsx não encontrado em nenhum dos caminhos "
                       "conhecidos (%s).", ", ".join(str(c) for c in _candidatos()))
        return {}
    try:
        mtime = os.path.getmtime(caminho)
    except OSError:
        logger.warning("TemplateCarteiras.xlsx inacessível em %s.", caminho)
        return {}

    with _lock:
        if _cache["caminho"] != caminho or _cache["mtime"] != mtime:
            try:
                _cache["carteiras"] = _ler_carteiras(caminho)
                _cache["caminho"] = caminho
                _cache["mtime"] = mtime
            except Exception:
                logger.exception(
                    "Falha lendo TemplateCarteiras.xlsx em %s — mantendo o último "
                    "resultado válido em cache.", caminho)
        return _cache["carteiras"]


def info_arquivo():
    """{caminho, mtime, origem} do arquivo em uso (para a tela mostrar a
    fonte). origem: "compartilhada" | "clone" (cópia local, pode estar
    desatualizada) | None (arquivo não encontrado nesta máquina)."""
    carteiras_template()
    caminho = _cache["caminho"]
    if not caminho:
        origem = None
    elif Path(caminho) == _CAMINHO_CLONE:
        origem = "clone"
    else:
        origem = "compartilhada"
    return {"caminho": caminho, "mtime": _cache["mtime"], "origem": origem}


def wallets_bloqueadas_para_publicacao():
    """Contexto:
    {walletId} de carteiras com "Deve Publicar" != Sim no TemplateCarteiras.xlsx
    do ControleCargas — usado por beehus_console.py::filter_groupings_by_publish_state
    para tirar da lista de "disponíveis para publicar" qualquer agrupamento
    que tenha uma dessas carteiras entre os membros.

    Fail-open: sem arquivo -> set() vazio (não bloqueia nada), mesma postura
    de antes — um dado auxiliar de outro projeto não derruba a tela.
    """
    return {wid for wid, info in carteiras_template().items() if not info["devePublicar"]}
